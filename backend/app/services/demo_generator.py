from pathlib import Path
from app.core.config import SAMPLE_DOCS_DIR
import pymupdf as fitz
import docx

OLD_POLICY_TEXT = """Apex University Academic & Examination Regulations (2025-2026 Edition)
[SYNTHETIC DEMONSTRATION POLICY DOCUMENT]

Section 1: Academic Attendance and Leave Regulations
1.1 Minimum Attendance Requirement: All registered undergraduate and postgraduate students must maintain at least 75% attendance in lectures and practical sessions to remain eligible for end-semester examinations.
1.2 Medical Leave Exemption: Students unable to meet the attendance threshold due to hospitalization may submit medical certificates to the Dean of Academic Affairs within 7 days of discharge.
1.3 Duty Leave Provision: Official athletic and co-curricular representation leaves shall be credited as attended instructional hours upon verification.

Section 2: Examination Enrollment and Fee Structure
2.1 Examination Registration Fee: Every candidate must remit an examination registration fee of $100 per semester prior to hall ticket issuance.
2.2 Re-Evaluation Requests: A candidate may apply for re-evaluation of answer scripts within 10 days of grade publication upon payment of $50 per course.
2.3 Hall Ticket Inspection: Students must visibly display their official institutional identity card when entering the university examination hall.

Section 3: Project Deadlines and Submission Guidelines
3.1 Capstone Submission Deadline: Final capstone project reports and portfolios must be submitted on or before 15 July 2025 to the departmental examination coordinator.
3.2 Late Submission Policy: Late submission of academic projects is permitted with an administrative late fine of $25 up to 5 business days past deadline.
3.3 Thesis Binding Specification: Physical hardcopy dissertations must be bound in maroon leatherette cloth prior to oral defense.

Section 4: Merit Scholarships and Academic Honors
4.1 Dean's Merit Scholarship: All enrolled full-time students in good academic standing are eligible to apply for institutional tuition scholarships.
4.2 Minimum Qualifying Grade: Scholarship awardees must maintain continuous academic progress with no unexcused semester backlogs.
"""

NEW_POLICY_TEXT = """Apex University Academic & Examination Regulations (2026-2027 Edition)
[SYNTHETIC DEMONSTRATION POLICY DOCUMENT]

Section 1: Academic Attendance and Leave Regulations
1.1 Minimum Attendance Requirement: All registered undergraduate and postgraduate students must maintain at least 80% attendance in lectures and practical sessions to remain eligible for end-semester examinations.
1.2 Medical Leave Exemption: Students unable to meet the attendance threshold due to hospitalization may submit medical certificates to the Dean of Academic Affairs within 7 days of discharge.
1.3 Biometric Morning Logging: Daily attendance must be registered using biometric turnstiles no later than 09:15 AM each instructional morning.

Section 2: Examination Enrollment and Fee Structure
2.1 Examination Registration Fee: Every candidate must remit an examination registration fee of $150 per semester prior to hall ticket issuance.
2.2 Re-Evaluation Requests: A candidate may apply for re-evaluation of answer scripts within 10 days of grade publication upon payment of $50 per course.
2.3 Hall Ticket Inspection: Valid institutional identity cards must be visibly carried upon entry to the university examination hall.

Section 3: Project Deadlines and Submission Guidelines
3.1 Capstone Submission Deadline: Final capstone project reports and portfolios must be submitted on or before 30 July 2026 to the departmental examination coordinator.
3.2 Late Submission Policy: Late submission of academic projects is strictly prohibited under any circumstances; late portals will not accept files.
3.3 Mandatory Ethics Requirement: All graduating candidates must successfully complete a mandatory 2-credit Artificial Intelligence & Academic Integrity seminar before degree clearance.

Section 4: Merit Scholarships and Academic Honors
4.1 Dean's Merit Scholarship: Only final-year undergraduate students with a cumulative GPA of 3.8 or above are eligible to apply for institutional tuition scholarships.
4.2 Minimum Qualifying Grade: Scholarship awardees must maintain continuous academic progress with no unexcused semester backlogs.
"""


def generate_sample_files():
    """Generates TXT, PDF, and DOCX files for both old and new synthetic policies."""
    SAMPLE_DOCS_DIR.mkdir(parents=True, exist_ok=True)

    # 1. Write TXT
    old_txt_path = SAMPLE_DOCS_DIR / "Apex_University_Academic_Policy_2025.txt"
    new_txt_path = SAMPLE_DOCS_DIR / "Apex_University_Academic_Policy_2026.txt"
    old_txt_path.write_text(OLD_POLICY_TEXT, encoding="utf-8")
    new_txt_path.write_text(NEW_POLICY_TEXT, encoding="utf-8")

    # 2. Write PDF using PyMuPDF
    old_pdf_path = SAMPLE_DOCS_DIR / "Apex_University_Academic_Policy_2025.pdf"
    new_pdf_path = SAMPLE_DOCS_DIR / "Apex_University_Academic_Policy_2026.pdf"

    for txt_content, pdf_path in [
        (OLD_POLICY_TEXT, old_pdf_path),
        (NEW_POLICY_TEXT, new_pdf_path),
    ]:
        doc = fitz.open()
        page = doc.new_page()
        # insert textbox
        rect = fitz.Rect(50, 50, 550, 750)
        page.insert_textbox(rect, txt_content, fontsize=10, fontname="helv")
        doc.save(str(pdf_path))
        doc.close()

    # 3. Write DOCX
    old_docx_path = SAMPLE_DOCS_DIR / "Apex_University_Academic_Policy_2025.docx"
    new_docx_path = SAMPLE_DOCS_DIR / "Apex_University_Academic_Policy_2026.docx"

    for txt_content, docx_path in [
        (OLD_POLICY_TEXT, old_docx_path),
        (NEW_POLICY_TEXT, new_docx_path),
    ]:
        d = docx.Document()
        for line in txt_content.split("\n"):
            line = line.strip()
            if not line:
                continue
            if line.startswith("Section ") or line.startswith("Apex University"):
                d.add_heading(line, level=1 if line.startswith("Apex") else 2)
            else:
                d.add_paragraph(line)
        d.save(str(docx_path))

    return {
        "old_txt": str(old_txt_path),
        "new_txt": str(new_txt_path),
        "old_pdf": str(old_pdf_path),
        "new_pdf": str(new_pdf_path),
        "old_docx": str(old_docx_path),
        "new_docx": str(new_docx_path),
    }
