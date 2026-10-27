import os
from pathlib import Path
from typing import Optional
from pydantic import BaseModel
import pymupdf as fitz
import docx


class ExtractedPage(BaseModel):
    page_number: int
    text: str


class ExtractedDocument(BaseModel):
    filename: str
    file_type: str
    full_text: str
    pages: list[ExtractedPage]
    total_characters: int
    is_scanned: bool = False
    warning: Optional[str] = None


class DocumentExtractionError(Exception):
    pass


class ScannedPdfError(DocumentExtractionError):
    pass


class DocumentExtractor:
    @staticmethod
    def extract_from_bytes(file_bytes: bytes, filename: str) -> ExtractedDocument:
        ext = Path(filename).suffix.lower()
        if ext == ".pdf":
            return DocumentExtractor._extract_pdf(file_bytes, filename)
        elif ext == ".docx":
            return DocumentExtractor._extract_docx(file_bytes, filename)
        elif ext == ".txt":
            return DocumentExtractor._extract_txt(file_bytes, filename)
        else:
            raise DocumentExtractionError(f"Unsupported file format '{ext}'. Allowed: .pdf, .docx, .txt")

    @staticmethod
    def _extract_pdf(file_bytes: bytes, filename: str) -> ExtractedDocument:
        pages: list[ExtractedPage] = []
        try:
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            total_chars = 0
            for page_idx in range(len(doc)):
                page = doc[page_idx]
                page_text = page.get_text("text") or ""
                cleaned = page_text.strip()
                total_chars += len(cleaned)
                pages.append(ExtractedPage(page_number=page_idx + 1, text=cleaned))
            doc.close()

            # Detect scanned PDF condition
            if len(pages) > 0 and total_chars < 30:
                raise ScannedPdfError(
                    f"Scanned PDF detected for '{filename}'. Total extracted text has only {total_chars} characters. "
                    "This document appears to contain scanned images without an embedded text layer; OCR is required."
                )

            full_text = "\n\n".join(p.text for p in pages if p.text)
            return ExtractedDocument(
                filename=filename,
                file_type="pdf",
                full_text=full_text,
                pages=pages,
                total_characters=total_chars,
                is_scanned=False,
            )
        except ScannedPdfError:
            raise
        except Exception as e:
            raise DocumentExtractionError(f"Failed to read PDF file '{filename}': {str(e)}")

    @staticmethod
    def _extract_docx(file_bytes: bytes, filename: str) -> ExtractedDocument:
        try:
            import io
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = []
            for p in doc.paragraphs:
                text = p.text.strip()
                if text:
                    paragraphs.append(text)

            # Check tables as well
            for table in doc.tables:
                for row in table.rows:
                    row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_texts:
                        paragraphs.append(" | ".join(row_texts))

            full_text = "\n\n".join(paragraphs)
            pages = [ExtractedPage(page_number=1, text=full_text)]
            return ExtractedDocument(
                filename=filename,
                file_type="docx",
                full_text=full_text,
                pages=pages,
                total_characters=len(full_text),
                is_scanned=False,
            )
        except Exception as e:
            raise DocumentExtractionError(f"Failed to parse DOCX file '{filename}': {str(e)}")

    @staticmethod
    def _extract_txt(file_bytes: bytes, filename: str) -> ExtractedDocument:
        try:
            # Try utf-8 first, fallback to latin-1
            try:
                text = file_bytes.decode("utf-8")
            except UnicodeDecodeError:
                text = file_bytes.decode("latin-1")

            cleaned = text.strip()
            pages = [ExtractedPage(page_number=1, text=cleaned)]
            return ExtractedDocument(
                filename=filename,
                file_type="txt",
                full_text=cleaned,
                pages=pages,
                total_characters=len(cleaned),
                is_scanned=False,
            )
        except Exception as e:
            raise DocumentExtractionError(f"Failed to read TXT file '{filename}': {str(e)}")
